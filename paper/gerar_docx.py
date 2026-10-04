# Gera o manuscrito .docx no formato da Revista Brasileira de Criminalística (template V2)
# a partir de paper/manuscrito.md. Produz duas versões, como exige a submissão duplo-cega:
#   paper/saida/manuscrito_RBC_identificado.docx
#   paper/saida/manuscrito_RBC_anonimo.docx
#
# Convenções do manuscrito.md:
#   bloco inicial entre linhas "---" com chaves: titulo, autores, afiliacoes (lista), email,
#   resumo, palavras_chave, abstract, keywords, agradecimentos
#   "# 1. TÍTULO"      -> título de seção (10 pt, negrito, maiúsculas)
#   "## 1.1. Subtítulo" -> subtítulo 1 (negrito itálico)
#   "### 1.1.1. Sub"    -> subtítulo 2 (itálico)
#   "![Legenda](figuras/x.png){largura=coluna|pagina}" -> figura + legenda (8 pt) depois
#   "Tabela N. legenda" seguida de tabela em pipes -> legenda (8 pt) antes da tabela
#   "# Referências Bibliográficas" -> título sem número; itens "[n] ..." em 10 pt
#   **negrito** e *itálico* inline; linhas em branco separam parágrafos.
import os, re, sys
from docx import Document
from docx.shared import Pt, Mm, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AQUI = os.path.dirname(os.path.abspath(__file__))
FONTE = "Times New Roman"
LARG_COLUNA_MM = 80   # (210 - 25 - 20 - 5) / 2
LARG_PAGINA_MM = 165


def ler_manuscrito(caminho):
    texto = open(caminho, encoding="utf-8").read()
    m = re.match(r"^---\n(.*?)\n---\n(.*)$", texto, re.S)
    meta_txt, corpo = m.group(1), m.group(2)
    meta, chave = {}, None
    for linha in meta_txt.split("\n"):
        if re.match(r"^[a-z_]+:", linha):
            chave, _, val = linha.partition(":")
            val = val.strip()
            meta[chave] = [] if val == "" else val
        elif linha.startswith("  - ") and chave:
            meta[chave].append(linha[4:].strip())
        elif linha.startswith("  ") and chave and isinstance(meta[chave], str):
            meta[chave] += " " + linha.strip()
    return meta, corpo


def run_inline(par, texto, tamanho, negrito=False, italico=False):
    """Adiciona runs ao parágrafo interpretando **negrito** e *itálico*."""
    for tok in re.split(r"(\*\*.+?\*\*|\*.+?\*|\^.+?\^)", texto):
        if not tok:
            continue
        b, i, t, sup = negrito, italico, tok, False
        if tok.startswith("**"):
            b, t = True, tok[2:-2]
        elif tok.startswith("*"):
            i, t = True, tok[1:-1]
        elif tok.startswith("^"):
            sup, t = True, tok[1:-1]
        r = par.add_run(t)
        r.font.superscript = sup
        r.font.name = FONTE; r._element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
        r.font.size = Pt(tamanho); r.bold = b; r.italic = i


def paragrafo(doc, texto, tamanho=10, alinh=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=0, depois=0,
              recuo=True, negrito=False, italico=False, espaco=1.0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = alinh; pf.space_before = Pt(antes); pf.space_after = Pt(depois)
    pf.line_spacing = espaco
    if recuo:
        pf.first_line_indent = Cm(0.5)
    run_inline(p, texto, tamanho, negrito, italico)
    return p


def set_colunas(secao, n, espaco_twips=284):
    sectPr = secao._sectPr
    for c in sectPr.xpath("./w:cols"):
        sectPr.remove(c)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(n)); cols.set(qn("w:space"), str(espaco_twips)); cols.set(qn("w:equalWidth"), "true")
    sectPr.append(cols)


def margens(secao):
    secao.page_width, secao.page_height = Mm(210), Mm(297)
    secao.top_margin, secao.bottom_margin = Mm(10), Mm(20)
    secao.left_margin, secao.right_margin = Mm(25), Mm(20)


def abrir_largura_pagina(doc):
    sec = doc.add_section(WD_SECTION.CONTINUOUS); margens(sec); set_colunas(sec, 1)


def fechar_largura_pagina(doc):
    sec = doc.add_section(WD_SECTION.CONTINUOUS); margens(sec); set_colunas(sec, 2)


def tabela_md(doc, linhas, largura_mm=LARG_COLUNA_MM):
    cab = [c.strip() for c in linhas[0].strip("|").split("|")]
    corpo = [[c.strip() for c in l.strip("|").split("|")] for l in linhas[2:]]
    t = doc.add_table(rows=1 + len(corpo), cols=len(cab))
    t.style = "Table Grid"; t.alignment = 1; t.autofit = False
    # primeira coluna (rótulos) mais larga; demais iguais
    n = len(cab); w0 = largura_mm * (0.34 if n > 2 else 0.5); wi = (largura_mm - w0) / max(n - 1, 1)
    larguras = [w0] + [wi] * (n - 1)
    for row in t.rows:
        for j, cel in enumerate(row.cells):
            cel.width = Mm(larguras[j])
    for j, c in enumerate(cab):
        cel = t.cell(0, j); cel.text = ""
        run_inline(cel.paragraphs[0], c, 8, negrito=True); cel.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, lin in enumerate(corpo, start=1):
        for j, c in enumerate(lin):
            cel = t.cell(i, j); cel.text = ""
            run_inline(cel.paragraphs[0], c, 8)
            cel.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
    for row in t.rows:
        for cel in row.cells:
            for p in cel.paragraphs:
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def construir(meta, corpo, anonimo, saida):
    doc = Document()
    est = doc.styles["Normal"]; est.font.name = FONTE; est.font.size = Pt(10)
    est.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)

    # --------- Seção 1 (coluna única): título, autores, resumo, abstract
    sec = doc.sections[0]; margens(sec); set_colunas(sec, 1)
    paragrafo(doc, "DOI: http://dx.doi.org/10.15260/rbc.v0i0.00", 8, WD_ALIGN_PARAGRAPH.LEFT, recuo=False)
    paragrafo(doc, meta["titulo"], 17, WD_ALIGN_PARAGRAPH.LEFT, antes=18, depois=12, recuo=False, negrito=True)
    if anonimo:
        paragrafo(doc, "[Autores omitidos para avaliação duplo-cega]", 13, WD_ALIGN_PARAGRAPH.LEFT, depois=8, recuo=False)
        paragrafo(doc, "[Filiações omitidas para avaliação duplo-cega]", 8, WD_ALIGN_PARAGRAPH.LEFT, recuo=False, italico=True)
    else:
        paragrafo(doc, meta["autores"], 13, WD_ALIGN_PARAGRAPH.LEFT, depois=8, recuo=False)
        for af in meta["afiliacoes"]:
            paragrafo(doc, af, 8, WD_ALIGN_PARAGRAPH.LEFT, recuo=False, italico=True)
        paragrafo(doc, f"*Endereço de e-mail para correspondência: {meta['email']}", 8, WD_ALIGN_PARAGRAPH.LEFT, recuo=False, italico=True)
    paragrafo(doc, "Recebido em 00/00/0000; Revisado em 00/00/0000; Aceito em 00/00/0000", 8, WD_ALIGN_PARAGRAPH.LEFT, antes=6, depois=20, recuo=False)

    paragrafo(doc, "Resumo", 9, WD_ALIGN_PARAGRAPH.LEFT, recuo=False, negrito=True)
    paragrafo(doc, meta["resumo"], 9, recuo=False, depois=4)
    paragrafo(doc, f"**Palavras-Chave:** {meta['palavras_chave']}", 8, WD_ALIGN_PARAGRAPH.LEFT, recuo=False, depois=10)
    paragrafo(doc, "Abstract", 9, WD_ALIGN_PARAGRAPH.LEFT, recuo=False, negrito=True)
    paragrafo(doc, meta["abstract"], 9, recuo=False, depois=4)
    paragrafo(doc, f"**Keywords:** {meta['keywords']}", 8, WD_ALIGN_PARAGRAPH.LEFT, recuo=False, depois=6)

    # --------- Seção 2 (duas colunas): corpo
    sec2 = doc.add_section(WD_SECTION.CONTINUOUS); margens(sec2); set_colunas(sec2, 2)

    linhas = corpo.split("\n"); i = 0; em_refs = False
    while i < len(linhas):
        l = linhas[i].rstrip()
        if not l.strip():
            i += 1; continue
        if l.startswith("# "):
            titulo = l[2:].strip()
            em_refs = titulo.lower().startswith("referências")
            texto = titulo if em_refs or titulo.lower().startswith("agradecimentos") else titulo.upper()
            paragrafo(doc, texto, 10, WD_ALIGN_PARAGRAPH.LEFT, antes=12, depois=6, recuo=False, negrito=True)
        elif l.startswith("## "):
            paragrafo(doc, l[3:].strip(), 10, WD_ALIGN_PARAGRAPH.LEFT, antes=8, depois=4, recuo=False, negrito=True, italico=True)
        elif l.startswith("### "):
            paragrafo(doc, l[4:].strip(), 10, WD_ALIGN_PARAGRAPH.LEFT, antes=6, depois=3, recuo=False, italico=True)
        elif l.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)(\{largura=(\w+)\})?", l)
            legenda, caminho, _, larg = m.groups()
            larg_mm = LARG_PAGINA_MM if larg == "pagina" else LARG_COLUNA_MM
            if larg == "pagina":
                abrir_largura_pagina(doc)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.add_run().add_picture(os.path.join(AQUI, caminho), width=Mm(larg_mm))
            paragrafo(doc, legenda, 8, WD_ALIGN_PARAGRAPH.JUSTIFY if len(legenda) > 90 else WD_ALIGN_PARAGRAPH.CENTER, depois=8, recuo=False)
            if larg == "pagina":
                fechar_largura_pagina(doc)
        elif re.match(r"^Tabela \d+\.", l):
            pagina = "{largura=pagina}" in l
            l = l.replace("{largura=pagina}", "").rstrip()
            if pagina:
                abrir_largura_pagina(doc)
            paragrafo(doc, l, 8, WD_ALIGN_PARAGRAPH.CENTER, antes=6, depois=3, recuo=False)
            i += 1
            while i < len(linhas) and not linhas[i].strip():
                i += 1
            bloco = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                bloco.append(linhas[i]); i += 1
            tabela_md(doc, bloco, LARG_PAGINA_MM if pagina else LARG_COLUNA_MM)
            if pagina:
                fechar_largura_pagina(doc)
            continue
        elif em_refs and re.match(r"^\[\d+\]", l):
            if anonimo and any(n in l for n in meta.get("ocultar_anonimo", [])):
                l = re.match(r"^\[\d+\]", l).group(0) + " [Referência omitida para avaliação duplo-cega.]"
            p = paragrafo(doc, l, 10, WD_ALIGN_PARAGRAPH.JUSTIFY, depois=3, recuo=False)
            p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.6)
        else:
            # parágrafo (pode continuar em linhas seguintes até linha vazia)
            bloco = [l]
            while i + 1 < len(linhas) and linhas[i + 1].strip() and not re.match(r"^(#|!\[|\||Tabela \d+\.)", linhas[i + 1]):
                i += 1; bloco.append(linhas[i].rstrip())
            texto = " ".join(bloco)
            if anonimo:
                for nome in meta.get("ocultar_anonimo", []):
                    texto = texto.replace(nome, "[omitido]")
            paragrafo(doc, texto, 10, depois=4)
        i += 1

    os.makedirs(os.path.dirname(saida), exist_ok=True)
    doc.save(saida); print("gerado", saida)


if __name__ == "__main__":
    meta, corpo = ler_manuscrito(os.path.join(AQUI, "manuscrito.md"))
    construir(meta, corpo, False, os.path.join(AQUI, "saida", "manuscrito_RBC_identificado.docx"))
    construir(meta, corpo, True, os.path.join(AQUI, "saida", "manuscrito_RBC_anonimo.docx"))
